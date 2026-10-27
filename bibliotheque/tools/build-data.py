#!/usr/bin/env python3
"""Prépare les données de la Bibliothèque des compétences à partir des fichiers sources.

Entrées (non versionnées, à fournir en arguments ou dans le dossier sources/) :
  --produits  Base_produits_2027_*.xlsx            (onglets PRODUITS, PRODUITS CERTIFIANTS, LOTS)
  --rome      ROME4_V62_*_Vue_Consolidee.xlsx       (onglets VUE_METIERS, COMPETENCES, MACRO_COMPETENCES, ...)

Sorties (chargeables par double-clic, sans serveur) :
  data/catalogue.js          window.LDC_CATALOGUE   produits, certifications, lots
  data/rome4-compact.json    version JSON du référentiel ROME 4.0 compact
  data/rome4-compact.js      window.LDC_ROME4       même contenu pour le navigateur
  content/content.bundle.js  window.LDC_CONTENT     agrégat de content/*.json

Règles :
  - aucun tarif, aucun nom de collaborateur dans les sorties ;
  - libellés et types ROME conservés tels quels ;
  - aucune valeur n'est inventée : champ absent = null ;
  - le contrôle d'intégrité (mapping ROME, codes produits) échoue si un identifiant n'existe pas.
"""
import argparse, glob, hashlib, json, os, sys, warnings
import openpyxl

warnings.filterwarnings('ignore')
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))


def find(pattern, given):
    if given:
        return given
    hits = sorted(glob.glob(os.path.join(ROOT, 'sources', pattern)))
    if not hits:
        sys.exit(f'Fichier source introuvable ({pattern}). Passez-le en argument.')
    return hits[-1]


def sha(path):
    h = hashlib.sha256()
    with open(path, 'rb') as f:
        for chunk in iter(lambda: f.read(1 << 20), b''):
            h.update(chunk)
    return h.hexdigest()[:16]


def s(v):
    """Chaîne nettoyée ou None."""
    if v is None:
        return None
    t = str(v).replace('\xa0', ' ').strip()
    return t if t and t not in ('-', '#REF!', '0') else None


def num(v):
    try:
        return round(float(v), 2)
    except (TypeError, ValueError):
        return None


def write_js(path, var, obj):
    os.makedirs(os.path.dirname(path), exist_ok=True)
    with open(path, 'w', encoding='utf-8') as f:
        f.write('/* Fichier généré par tools/build-data.py : ne pas modifier à la main. */\n')
        f.write(f'window.{var}=' + json.dumps(obj, ensure_ascii=False, separators=(',', ':')) + ';\n')
    print(f'  {os.path.relpath(path, ROOT)}  {os.path.getsize(path) // 1024} Ko')


# --------------------------------------------------------------------------------------------
def build_catalogue(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    # PRODUITS
    rows = list(wb['PRODUITS'].iter_rows(values_only=True))
    H = {h: i for i, h in enumerate(rows[0]) if h}
    produits, seen = [], set()
    for r in rows[3:]:
        code = s(r[H['CODE FORMEO']])
        if not code or code in seen:
            continue
        seen.add(code)
        niveau = s(r[H['NIVEAU']])
        produits.append({
            'c': code,
            'ob': s(r[H['CODE ONEBASE']]),
            't': s(r[H['LIBELLE PRODUIT']]),
            'pl': s(r[H['PRODUIT/LOT']]),
            'tp': s(r[H['TYPE PRODUIT/LOT']]),
            'f': s(r[H['FORMAT']]),
            'n': niveau,
            'vs': s(r[H['VENDU SEUL']]),
            'dec': s(r[H['DECISION']]),
            'mq': s(r[H['MARQUE']]),
            'm': s(r[H['MARCHE FORMEO']]),
            'dom': s(r[H['DOMAINE FORMEO']]),
            'th': s(r[H['THEME FORMEO']]),
            'dl': s(r[H['DUREE LITTERAIRE']]),
            'dh': s(r[H['DUREE LITTERAIRE EN HEURES']]),
            'dj': num(r[H['DUREE EN JOURS']]),
            'best': 1 if s(r[H['BEST']]) == 'OUI' else 0,
            'cpf': 1 if s(r[H['ELIGIBLE CPF']]) == 'OUI' else 0,
        })
    # PRODUITS CERTIFIANTS (référentiel France compétences : RNCP / RS)
    rows = list(wb['PRODUITS CERTIFIANTS'].iter_rows(values_only=True))
    CH = {h: i for i, h in enumerate(rows[0]) if h}
    certifs = {}
    for r in rows[1:]:
        code = s(r[CH['CODE FORMEO']])
        if not code:
            continue
        ech = r[CH['ECHEANCE CERTIFICATION']]
        certifs[code] = {
            'rsRncp': s(r[CH['CODE RS/RNCP']]),
            'bpf': s(r[CH['NIVEAU BPF']]),
            'modalite': s(r[CH['MODALITE EVALUATION']]),
            'categorie': s(r[CH['CATEGORIE PARCOURS']]),
            'partenariat': s(r[CH['PARTENARIAT']]),
            'statut': s(r[CH['STATUT CERTIFICATION']]),
            'echeance': ech.strftime('%Y-%m-%d') if hasattr(ech, 'strftime') else s(ech),
            'urlLdc': s(r[CH['URL SITE LDC']]),
            'urlFc': s(r[CH['URL FC']]),
        }
    # LOTS : composition des packs et parcours (source des « produits complémentaires » factuels)
    ws = wb['LOTS']
    lots, blank = {}, 0
    for i, r in enumerate(ws.iter_rows(values_only=True)):
        if i < 3:
            continue
        if r[0] is None:
            blank += 1
            if blank > 50:
                break
            continue
        blank = 0
        lot, rang, comp = s(r[0]), s(r[1]), s(r[2])
        if not lot or comp is None:
            continue
        rang = int(rang) if rang and rang.isdigit() else 0
        entry = lots.setdefault(lot, {'membres': []})
        if rang == 0 and comp == lot:
            entry['t'] = s(r[3]); entry['f'] = s(r[6])
        elif rang > 0:
            entry['membres'].append(comp)
    lots = {k: {'t': v.get('t'), 'f': v.get('f'), 'membres': v['membres']} for k, v in lots.items() if v['membres']}
    return produits, certifs, lots


# --------------------------------------------------------------------------------------------
def build_rome(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)

    def sheet(name):
        rows = wb[name].iter_rows(values_only=True)
        hdr = next(rows)
        idx = {h: i for i, h in enumerate(hdr)}
        return idx, rows

    readme = {}
    for r in wb['README'].iter_rows(values_only=True):
        if r[0] in ('Version', 'Validation', 'Publication'):
            readme[r[0]] = s(r[1])
    # hiérarchie de compétences + compétences
    idx, rows = sheet('COMPETENCES')
    types, natures, sous = [], [], []
    def pos(lst, v):
        if v not in lst:
            lst.append(v)
        return lst.index(v)
    domc, enjeux, objectifs, comp = {}, {}, {}, []
    for r in rows:
        cid = r[idx['id_competence']]
        if cid is None:
            continue
        domc[str(r[idx['code_domaine_competence']])] = r[idx['domaine_competence']]
        enjeux[str(r[idx['code_enjeu']])] = (r[idx['enjeu']], str(r[idx['code_domaine_competence']]))
        objectifs[str(r[idx['code_objectif']])] = (r[idx['objectif']], str(r[idx['code_enjeu']]))
        comp.append([int(cid), r[idx['libelle_competence']], pos(types, r[idx['type_rome']]), pos(natures, r[idx['nature']]),
                     pos(sous, r[idx['sous_cat_comp']]), str(r[idx['code_macro_competence']])])
    # macro-compétences + savoir-être professionnels
    idx, rows = sheet('MACRO_COMPETENCES')
    macro = []
    for r in rows:
        if r[idx['code_macro_comp']] is None:
            continue
        macro.append([str(r[idx['code_macro_comp']]), int(r[idx['id_macro_competence']]), r[idx['libelle_macro_competence']],
                      pos(types, r[idx['type_rome']]), str(r[idx['code_objectif']])])
    idx, rows = sheet('SAVOIR_ETRE_PRO')
    sep = [[int(r[idx['id_macro_competence']]), str(r[idx['code_macro_comp']]), r[idx['libelle_macro_competence']]] for r in rows if r[idx['code_macro_comp']]]
    # métiers
    idx, rows = sheet('VUE_METIERS')
    gd, dp, metiers = {}, {}, []
    for r in rows:
        code = s(r[idx['code_rome']])
        if not code:
            continue
        gd[str(r[idx['code_grand_domaine']])] = r[idx['libelle_grand_domaine']]
        dp[str(r[idx['code_domaine_professionnel']])] = (r[idx['libelle_domaine_professionnel']], str(r[idx['code_grand_domaine']]))
        metiers.append([code, r[idx['libelle_rome']], str(r[idx['code_grand_domaine']]), str(r[idx['code_domaine_professionnel']])])
    return {
        'meta': {'version': readme.get('Version'), 'validation': readme.get('Validation'), 'publication': readme.get('Publication'),
                 'regle': 'Libellés, types et identifiants officiels, conservés tels quels. Aucune compétence ajoutée.',
                 'nbCompetences': len(comp), 'nbMacro': len(macro), 'nbMetiers': len(metiers)},
        'types': types, 'natures': natures, 'sousCategories': sous,
        'domainesCompetence': [[k, v] for k, v in sorted(domc.items())],
        'enjeux': [[k, v[0], v[1]] for k, v in sorted(enjeux.items())],
        'objectifs': [[k, v[0], v[1]] for k, v in sorted(objectifs.items())],
        'macro': macro, 'savoirEtrePro': sep, 'competences': comp,
        'grandsDomaines': [[k, v] for k, v in sorted(gd.items())],
        'domainesPro': [[k, v[0], v[1]] for k, v in sorted(dp.items())],
        'metiers': metiers,
    }


# --------------------------------------------------------------------------------------------
def bundle_content():
    out = {}
    cdir = os.path.join(ROOT, 'content')
    for fn in sorted(glob.glob(os.path.join(cdir, '*.json'))):
        key = os.path.splitext(os.path.basename(fn))[0]
        with open(fn, encoding='utf-8') as f:
            out[key] = json.load(f)
    # la vidéo de transition n'est déclarée disponible que si les fichiers existent (évite toute requête vers un fichier absent)
    br = out.get('brand', {})
    for k in ('logo', 'symbole'):
        if br.get(k):
            br[k]['disponible'] = os.path.exists(os.path.join(ROOT, br[k]['svg']))
    tr = out.get('transition', {}).get('video')
    if tr:
        tr['disponible'] = {k: os.path.exists(os.path.join(ROOT, tr[k])) for k in ('webm', 'mp4', 'poster') if tr.get(k)}
    write_js(os.path.join(cdir, 'content.bundle.js'), 'LDC_CONTENT', out)
    return out


def validate(content, produits, rome):
    """Contrôle d'intégrité : chaque identifiant cité dans content/ doit exister avec le libellé exact."""
    errs = []
    codes = {p['c'] for p in produits}
    comp = {c[0]: c for c in rome['competences']}
    mp = content.get('mapping-produits-rome', {}).get('produits', {})
    for code, entry in mp.items():
        if code not in codes:
            errs.append(f'mapping: code produit inconnu {code}')
        for c in entry.get('competences', []):
            ref = comp.get(int(c['id']))
            if not ref:
                errs.append(f'mapping {code}: id ROME inexistant {c["id"]}')
            elif c.get('libelle') and c['libelle'] != ref[1]:
                errs.append(f'mapping {code}: libellé différent pour {c["id"]} ({c["libelle"]!r} ≠ {ref[1]!r})')
    for code, entry in mp.items():
        if entry.get('statut') == 'VALIDE' and not (entry.get('validePar') and entry.get('date')):
            errs.append(f'mapping {code}: statut VALIDE sans « validePar » et « date »')
    metiers = {m[0] for m in rome['metiers']}
    for code, entry in mp.items():
        for mc in entry.get('metiers', []):
            if mc not in metiers:
                errs.append(f'mapping {code}: code métier ROME inexistant {mc}')
    doms = {p['dom'] for p in produits if p.get('dom')}
    for mo in content.get('domaines', {}).get('mondes', []):
        for d in mo.get('domaines', []):
            if d not in doms:
                errs.append(f'domaines.json: domaine FORMEO inconnu dans le catalogue « {d} »')
    rattaches = {d for mo in content.get('domaines', {}).get('mondes', []) for d in mo.get('domaines', [])}
    orphelins = sorted(doms - rattaches)
    if orphelins:
        print('  Avertissement : domaines du catalogue sans univers :', ', '.join(orphelins))
    for code in content.get('produits-commercial', {}).get('produits', {}):
        if code not in codes:
            errs.append(f'produits-commercial: code produit inconnu {code}')
    for m in content.get('moments-de-vie', {}).get('moments', []):
        for cid in m.get('competences', []):
            if int(cid) not in comp:
                errs.append(f'moment {m.get("id")}: id ROME inexistant {cid}')
    return errs


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--produits'); ap.add_argument('--rome')
    a = ap.parse_args()
    pp, rp = find('Base_produits_2027*.xlsx', a.produits), find('ROME4_*Vue_Consolidee*.xlsx', a.rome)
    print('Catalogue :', os.path.basename(pp)); print('ROME 4.0  :', os.path.basename(rp))
    produits, certifs, lots = build_catalogue(pp)
    rome = build_rome(rp)
    write_js(os.path.join(ROOT, 'data', 'catalogue.js'), 'LDC_CATALOGUE', {
        'meta': {'source': os.path.basename(pp), 'empreinte': sha(pp), 'nbProduits': len(produits),
                 'avertissement': 'Tarifs et noms de collaborateurs volontairement exclus.'},
        'produits': produits, 'certifications': certifs, 'lots': lots})
    jp = os.path.join(ROOT, 'data', 'rome4-compact.json')
    with open(jp, 'w', encoding='utf-8') as f:
        json.dump(rome, f, ensure_ascii=False, separators=(',', ':'))
    print(f'  data/rome4-compact.json  {os.path.getsize(jp) // 1024} Ko')
    write_js(os.path.join(ROOT, 'data', 'rome4-compact.js'), 'LDC_ROME4', rome)
    content = bundle_content()
    errs = validate(content, produits, rome)
    if errs:
        print('\nERREURS D\'INTÉGRITÉ :'); [print('  -', e) for e in errs]; sys.exit(1)
    print('\nIntégrité OK :', len(produits), 'produits,', len(certifs), 'certifications,', len(lots), 'lots,', len(rome['competences']), 'compétences ROME 4.')


if __name__ == '__main__':
    main()
